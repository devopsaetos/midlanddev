# -*- coding: utf-8 -*-
"""Company-wise Excel template + import for Maintenance.

Workbook layout (made by "Download Template"):
  Houses       one row per house: society, plot, file, member, nominee, size, meter ...
  Bill History one row per house + charge type + month: bill amount, paid, balance
  Lists        allowed values for the drop-downs
  Instructions how to fill it

Import matches a house by Society + Plot No, a member by Tracking ID (MRN) or CNIC, and
a bill row by house + charge type + month, so the same workbook can be imported again:
it only adds what is new and updates what changed. "Check File" shows exactly what
Import will do without writing anything.
"""
import base64
import collections
import datetime
import io
import re

from markupsafe import Markup, escape

from odoo import _, api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - listed in external_dependencies
    openpyxl = None

from odoo.addons.real_estate.models.res_member import MOBILE_RE, _format_cnic, _format_mobile

HOUSES = 'Houses'
HISTORY = 'Bill History'
LISTS = 'Lists'

# (key, header, width, list name for the drop-down)
HOUSE_COLS = [
    ('company', 'Company', 18, None),
    ('society', 'Society*', 18, 'Society'),
    ('phase', 'Phase', 18, 'Phase'),
    ('sector', 'Sector', 12, 'Sector'),
    ('plot', 'Plot No*', 12, None),
    ('file_no', 'File No', 12, None),
    ('tracking', 'Tracking ID (MRN)', 16, None),
    ('member', 'Member Name*', 28, None),
    ('cnic', 'CNIC', 17, None),
    ('mobile', 'Mobile', 14, None),
    ('phone', 'Phone', 14, None),
    ('email', 'Email', 22, None),
    ('address', 'Address', 28, None),
    ('city', 'City', 12, None),
    ('kin_name', 'Nominee Name', 22, None),
    ('kin_cnic', 'Nominee CNIC', 17, None),
    ('kin_mobile', 'Nominee Mobile', 14, None),
    ('kin_relation', 'Nominee Relation', 14, 'Relation'),
    ('category', 'Category', 13, 'Category'),
    ('product', 'Product (Size)', 13, 'Product'),
    ('unit_class', 'Unit Class', 11, 'Unit Class'),
    ('area', 'Area (sqft)', 10, None),
    ('booking_date', 'Booking Date', 12, None),
    ('ownership', 'Ownership Type', 13, 'Ownership'),
    ('agent', 'Recovery Agent', 16, 'Recovery Agent'),
    ('exemption', 'Exemption', 10, 'Yes/No'),
    ('service_charge', 'Service Charge', 10, 'Yes/No'),
    ('meter_no', 'Meter No', 12, None),
    ('meter_reading', 'Opening Meter Reading', 12, None),
    ('notes', 'Notes', 28, None),
]
HISTORY_COLS = [
    ('society', 'Society*', 18, 'Society'),
    ('plot', 'Plot No*', 12, None),
    ('member', 'Member Name', 28, None),
    ('charge_type', 'Charge Type*', 12, 'Charge Type'),
    ('month', 'Month*', 11, None),
    ('amount', 'Bill Amount*', 12, None),
    ('paid', 'Paid Amount', 12, None),
    ('payment_date', 'Payment Date', 12, None),
    ('balance', 'Balance', 12, None),
    ('status', 'Status', 12, None),
    ('invoice', 'Invoice', 22, None),
]
CHARGE_TYPES = {'utility': 'utility', 'maintenance': 'utility', 'electricity': 'electricity'}
CHARGE_LABEL = {'utility': 'Utility', 'electricity': 'Electricity'}
STATUS_LABEL = {'paid': 'Paid', 'partial': 'Partially Paid', 'not_paid': 'Not Paid', 'in_payment': 'In Payment'}
YES = {'yes', 'y', 'true', '1', 'x'}
NO = {'no', 'n', 'false', '0'}
MONTHS = {m: i for i, m in enumerate(
    ['jan', 'feb', 'mar', 'apr', 'may', 'jun', 'jul', 'aug', 'sep', 'oct', 'nov', 'dec'], 1)}


def _norm(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip().lower()


def _name_key(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime('%Y-%m-%d')
    return re.sub(r'\s+', ' ', str(value)).strip()


def _number(value):
    """float, None for an empty cell, False when it is not a number."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(',', '').strip())
    except ValueError:
        return False


def _date(value):
    """date, None for an empty cell, False when it cannot be read."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # Excel serial number
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(value))
    return _date_text(str(value).strip()) or _month(str(value)) or False


def _date_text(text):
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%d.%m.%Y', '%Y/%m/%d', '%d-%b-%Y', '%d %b %Y',
                '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _month(value):
    """First day of the month, None for an empty cell, False when it cannot be read.
    Takes a date, 'Sep-2026', 'Sep-26', 'September 2026', '2026-09', '09/2026'."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (datetime.datetime, datetime.date)) or isinstance(value, (int, float)):
        d = _date(value)
        return d.replace(day=1) if d else False
    text = str(value).strip().replace('_', '-').replace('/', '-').replace(' ', '-').replace('.', '-')
    m = re.match(r'^([A-Za-z]{3})[A-Za-z]*-(\d{2}|\d{4})$', text)
    if m and m.group(1).lower() in MONTHS:
        year = int(m.group(2))
        return datetime.date(year + 2000 if year < 100 else year, MONTHS[m.group(1).lower()], 1)
    m = re.match(r'^(\d{4})-(\d{1,2})$', text) or re.match(r'^(\d{1,2})-(\d{4})$', text)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        year, month = (a, b) if a > 12 else (b, a)
        if 1 <= month <= 12:
            return datetime.date(year, month, 1)
    d = _date_text(str(value).strip())
    return d.replace(day=1) if d else False


def _same(current, new, field=None):
    """Whether an Excel value equals what Odoo has, ignoring layout (spaces, dashes, case)."""
    if isinstance(new, str) or isinstance(current, str):
        a, b = str(current or ''), str(new or '')
        if field in (None, 'cnic', 'kin_cnic', 'mobile', 'kin_mobile', 'phone'):
            da, db = re.sub(r'\D', '', a), re.sub(r'\D', '', b)
            if da and (da == db or field in (None, 'mobile') and _mobile(a) and _mobile(a) == _mobile(b)):
                return True
        return _norm(a) == _norm(b)
    return (current or False) == (new or False)


def _mobile(value):
    """Mobile in the 0320-7834234 layout the member form wants, or '' when it is not one."""
    digits = re.sub(r'\D', '', value or '')
    if digits.startswith('92') and len(digits) == 12:
        digits = '0' + digits[2:]
    elif len(digits) == 10 and digits.startswith('3'):
        digits = '0' + digits
    formatted = _format_mobile(digits)
    return formatted if formatted and MOBILE_RE.match(formatted) else ''


class MaintenanceExcelImport(models.TransientModel):
    _name = 'maintenance.excel.import'
    _description = 'Maintenance Excel Import'

    company_id = fields.Many2one('res.company', required=True, default=lambda self: self.env.company,
                                 domain=lambda self: [('id', 'in', self.env.user.company_ids.ids)])
    society_ids = fields.Many2many('society', string='Societies', domain="[('is_society', '=', True)]",
                                  compute='_compute_society_ids', store=True, readonly=False,
                                  help='Societies of the company. Houses of other societies are refused.')
    with_data = fields.Boolean('Fill With Current Data', default=True,
                               help='Download the template with every house and bill row already in Odoo. '
                                    'Untick for an empty template.')
    template_file = fields.Binary(readonly=True, attachment=False)
    template_name = fields.Char()
    import_file = fields.Binary('Excel File', attachment=False)
    import_name = fields.Char()
    state = fields.Selection([('draft', 'Upload'), ('checked', 'Checked'), ('done', 'Imported')], default='draft')
    result_html = fields.Html(readonly=True, sanitize=False)
    error_count = fields.Integer(readonly=True)
    file_ids = fields.Many2many('file', string='Imported Houses', readonly=True)

    @api.depends('company_id')
    def _compute_society_ids(self):
        for rec in self:
            rec.society_ids = self.env['society'].search(
                [('is_society', '=', True), ('company_id', '=', rec.company_id.id)])

    @api.onchange('import_file')
    def _onchange_import_file(self):
        self.state = 'draft'
        self.result_html = False
        self.error_count = 0

    # ------------------------------------------------------------------ lists

    def _lists(self):
        """Allowed values per drop-down: {list name: {normalised name: record or value}}."""
        env = self.env
        societies = self.society_ids
        phases = env['society'].search([('is_society', '=', False), ('society_id', 'in', societies.ids)])
        sectors = env['sector'].search([('society_id', 'in', societies.ids)])
        agents = env['res.users'].search([('maintenance_recovery_agent', '=', True)])
        relation = dict(env['res.member']._fields['kin_member_relation']._description_selection(env))
        ownership = dict(env['file']._fields['file_ownership_type']._description_selection(env))

        def by_name(records):
            out = {}
            for r in records.sorted('id'):
                out.setdefault(_norm(r.display_name if r._name == 'res.users' else r.name), r)
            return out

        return collections.OrderedDict([
            ('Society', by_name(societies)),
            ('Phase', by_name(phases)),
            ('Sector', by_name(sectors)),
            ('Category', by_name(env['plot.category'].search([]))),
            ('Product', by_name(env['unit.category.type'].search([]))),
            ('Unit Class', by_name(env['unit.class'].search([]))),
            ('Ownership', {_norm(v): k for k, v in ownership.items()}),
            ('Relation', {_norm(v): k for k, v in relation.items()}),
            ('Recovery Agent', by_name(agents)),
            ('Charge Type', {'utility': 'utility', 'electricity': 'electricity'}),
            ('Yes/No', {'yes': True, 'no': False}),
        ]), phases, sectors

    def _houses(self):
        return self.env['file'].with_context(active_test=False).search(
            [('society_id', 'in', self.society_ids.ids), ('unit_number', '!=', False),
             ('file_status', 'not in', ('cancel', 'merged_and_cancel'))], order='society_id, unit_number, id')

    # --------------------------------------------------------------- template

    def action_download_template(self):
        self.ensure_one()
        if not openpyxl:
            raise UserError(_('The Python library openpyxl is missing on the server.'))
        if not self.society_ids:
            raise UserError(_('%s has no society. Pick the societies to put in the template.', self.company_id.name))
        lists, phases, _sectors = self._lists()
        wb = openpyxl.Workbook()
        head_fill = PatternFill('solid', fgColor='1F4E78')
        need_fill = PatternFill('solid', fgColor='C00000')
        head_font = Font(bold=True, color='FFFFFF')

        # Lists sheet: one column per drop-down, readable labels
        ws_l = wb.active
        ws_l.title = LISTS
        list_ranges = {}
        labels = {
            'Ownership': [v for k, v in self.env['file']._fields['file_ownership_type']._description_selection(self.env)],
            'Relation': [v for k, v in self.env['res.member']._fields['kin_member_relation']._description_selection(self.env)],
            'Charge Type': ['Utility', 'Electricity'],
            'Yes/No': ['Yes', 'No'],
        }
        for col, (name, values) in enumerate(lists.items(), 1):
            items = labels.get(name) or sorted({r.display_name if r._name == 'res.users' else r.name
                                                for r in values.values()})
            cell = ws_l.cell(1, col, name)
            cell.fill, cell.font = head_fill, head_font
            for row, item in enumerate(items, 2):
                ws_l.cell(row, col, item)
            letter = get_column_letter(col)
            ws_l.column_dimensions[letter].width = 22
            list_ranges[name] = "'%s'!$%s$2:$%s$%s" % (LISTS, letter, letter, max(len(items) + 1, 2))

        def sheet(title, cols, rows, date_cols=(), month_cols=(), money_cols=()):
            ws = wb.create_sheet(title, index=len(wb.sheetnames) - 1)
            for c, (key, header, width, _l) in enumerate(cols, 1):
                cell = ws.cell(1, c, header)
                cell.font = head_font
                cell.fill = need_fill if header.endswith('*') else head_fill
                cell.alignment = Alignment(wrap_text=True, vertical='center')
                ws.column_dimensions[get_column_letter(c)].width = width
            ws.row_dimensions[1].height = 32
            ws.freeze_panes = 'A2'
            for r, values in enumerate(rows, 2):
                for c, (key, *_x) in enumerate(cols, 1):
                    value = values.get(key)
                    if value in (None, False, ''):
                        continue
                    cell = ws.cell(r, c, value)
                    if key in date_cols:
                        cell.number_format = 'DD-MM-YYYY'
                    elif key in month_cols:
                        cell.number_format = 'MMM-YYYY'
                    elif key in money_cols:
                        cell.number_format = '#,##0'
            last = max(len(rows) + 1, 2) + 3000
            for c, (key, header, width, list_name) in enumerate(cols, 1):
                letter = get_column_letter(c)
                if list_name:
                    dv = DataValidation(type='list', formula1=list_ranges[list_name], allow_blank=True,
                                        showErrorMessage=True, errorTitle=header,
                                        error=_('Pick a value from the list (sheet "%s").', LISTS))
                    ws.add_data_validation(dv)
                    dv.add('%s2:%s%s' % (letter, letter, last))
                elif key in date_cols or key in month_cols:
                    for r in range(len(rows) + 2, last + 1):
                        ws.cell(r, c).number_format = 'MMM-YYYY' if key in month_cols else 'DD-MM-YYYY'
            ws.auto_filter.ref = 'A1:%s%s' % (get_column_letter(len(cols)), max(len(rows) + 1, 2))
            return ws

        house_rows, history_rows = [], []
        if self.with_data:
            houses = self._houses()
            relation = dict(self.env['res.member']._fields['kin_member_relation']._description_selection(self.env))
            ownership = dict(self.env['file']._fields['file_ownership_type']._description_selection(self.env))
            for f in houses:
                m = f.membership_id
                house_rows.append({
                    'company': f.society_id.company_id.name, 'society': f.society_id.name,
                    'phase': f.phase_id.name, 'sector': f.sector_id.name, 'plot': f.unit_number,
                    'file_no': f.name, 'tracking': m.ref or f.tracking_id, 'member': m.name,
                    'cnic': m.cnic, 'mobile': m.mobile, 'phone': m.phone, 'email': m.email,
                    'address': m.street, 'city': m.city,
                    'kin_name': m.kin_name, 'kin_cnic': m.kin_cnic, 'kin_mobile': m.kin_mobile,
                    'kin_relation': relation.get(m.kin_member_relation), 'category': f.category_id.name,
                    'product': f.unit_category_type_id.name, 'unit_class': f.unit_class_id.name,
                    'area': f.standard_area or None, 'booking_date': f.booking_date,
                    'ownership': ownership.get(f.file_ownership_type),
                    'agent': f.maintenance_recovery_agent_id.name,
                    'exemption': 'Yes' if f.exemption else 'No',
                    'service_charge': 'Yes' if f.service_charge else 'No',
                    'meter_no': f.meter_no, 'meter_reading': f.opening_meter_reading or None, 'notes': f.notes,
                })
            history = self.env['maintenance.charges.history'].search(
                [('file_id', 'in', houses.ids)]).sorted(
                lambda h: (h.file_id.society_id.name or '', h.file_id.unit_number or '', h.charge_type, h.date, h.id))
            for h in history:
                history_rows.append({
                    'society': h.file_id.society_id.name, 'plot': h.file_id.unit_number,
                    'member': h.file_id.membership_id.name, 'charge_type': CHARGE_LABEL.get(h.charge_type),
                    'month': h.date, 'amount': h.amount, 'paid': h.amount_paid, 'payment_date': h.payment_date,
                    'balance': h.residual, 'status': STATUS_LABEL.get(h.payment_status, h.payment_status),
                    'invoice': h.invoice_id.name,
                })
        sheet(HOUSES, HOUSE_COLS, house_rows, date_cols=('booking_date',))
        sheet(HISTORY, HISTORY_COLS, history_rows, date_cols=('payment_date',), month_cols=('month',),
              money_cols=('amount', 'paid', 'balance'))
        self._instructions_sheet(wb, phases, head_font, head_fill)
        wb.active = 0

        out = io.BytesIO()
        wb.save(out)
        name = 'Maintenance Import - %s%s.xlsx' % (
            self.company_id.name.replace('/', '-'), '' if self.with_data else ' (empty)')
        self.write({'template_file': base64.b64encode(out.getvalue()), 'template_name': name})
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s/%s/template_file/%s?download=true' % (self._name, self.id, name),
            'target': 'self',
        }

    def _instructions_sheet(self, wb, phases, head_font, head_fill):
        ws = wb.create_sheet('Instructions')
        lines = [
            ('Maintenance Excel Import - %s' % self.company_id.name, True),
            ('Societies: %s' % ', '.join(self.society_ids.mapped('name')), False),
            ('Phases: %s' % ', '.join(phases.mapped('name')), False),
            ('', False),
            ('Sheet "Houses" - one row per house (plot)', True),
            ('Red columns (*) are required: Society, Plot No, Member Name.', False),
            ('A house is found by Society + Plot No. A new plot number makes a new file (File No is given by Odoo).', False),
            ('The member is found by Tracking ID (MRN), else by CNIC + name. Leave Tracking ID empty for a new member: Odoo gives the MRN.', False),
            ('Empty cells never clear data in Odoo: only filled cells are written.', False),
            ('Mobile: 0300-1234567 (also +92 300 1234567 is accepted). CNIC: 35202-1234567-1.', False),
            ('Dates: 22-07-2022. Exemption / Service Charge: Yes or No. Drop-down columns take values from sheet "Lists".', False),
            ('Opening Meter Reading = last reading before the first Monthly Bill (the first bill starts from it).', False),
            ('', False),
            ('Sheet "Bill History" - one row per house + charge type + month', True),
            ('Required: Society, Plot No, Charge Type (Utility / Electricity), Month (e.g. Sep-2026), Bill Amount.', False),
            ('Bill Amount = total payable of that month, previous balance included (as on the paper bill).', False),
            ('Paid Amount = what was paid against it. Balance = left to pay; when empty it is Bill Amount - Paid Amount.', False),
            ('The Balance of the latest month is the arrears that the next Monthly Bill shows.', False),
            ('Status, Member Name and Invoice are for reading only. Rows that have an Invoice come from real bills and are never changed.', False),
            ('The same month again for the same house + charge type updates that row (no duplicates).', False),
            ('', False),
            ('How to import', True),
            ('Maintenance Charges > Monthly Bills > Import from Excel: pick the company, upload the file, press Check File,', False),
            ('read the report (nothing is saved yet), then press Import. Rows with errors are skipped and listed.', False),
            ('', False),
            ('Example Houses row', True),
            ('Canal Valley-I | Canal Valley Ph-1 | Sector-1 | CV-R-103 | | MRN-000652 | MRS. SAIMA ASAD QURESHI | 35202-8098250-7 | 0300-8423022 | ... | 5 Marla | Plot | 1125 | 22-07-2022 | End User | | No | No | M-1234 | 860', False),
            ('Example Bill History rows', True),
            ('Canal Valley-I | CV-R-103 | | Utility | Aug-2026 | 2300 | 2300 | 10-08-2026 | 0', False),
            ('Canal Valley-I | CV-R-103 | | Utility | Sep-2026 | 2300 | 0 | | 2300', False),
        ]
        ws.column_dimensions['A'].width = 140
        for r, (text, bold) in enumerate(lines, 1):
            cell = ws.cell(r, 1, text)
            if bold:
                cell.font = Font(bold=True, size=12)

    # ------------------------------------------------------------------ check

    def _read_sheet(self, wb, title, cols):
        """[(excel row number, {key: value})] for the non-empty rows of a sheet."""
        ws = next((s for s in wb.worksheets if _norm(s.title) == _norm(title)), None)
        if ws is None:
            return None, []
        by_header = {_norm(h.rstrip('*')): k for k, h, _w, _l in cols}
        rows = ws.iter_rows(values_only=True)
        header = next(rows, ())
        index = {}
        for i, h in enumerate(header):
            key = by_header.get(_norm(str(h or '').rstrip('*')))
            if key and key not in index:
                index[key] = i
        out = []
        for n, row in enumerate(rows, 2):
            values = {k: row[i] if i < len(row) else None for k, i in index.items()}
            if all(_text(v) == '' for v in values.values()):
                continue
            out.append((n, values))
        return index, out

    def _plan(self):
        """Work out everything Import will do, without writing. Returns a plan dict."""
        self.ensure_one()
        if not self.import_file:
            raise UserError(_('Upload the Excel file first.'))
        if not openpyxl:
            raise UserError(_('The Python library openpyxl is missing on the server.'))
        try:
            wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.import_file)), data_only=True, read_only=True)
        except Exception as e:
            raise UserError(_('This is not an Excel (.xlsx) file: %s', e))
        house_index, house_rows = self._read_sheet(wb, HOUSES, HOUSE_COLS)
        history_index, history_rows = self._read_sheet(wb, HISTORY, HISTORY_COLS)
        if house_index is None and history_index is None:
            raise UserError(_('The file has no "%s" or "%s" sheet. Use Download Template to get the format.',
                              HOUSES, HISTORY))
        for title, index, needed in ((HOUSES, house_index, ('society', 'plot', 'member')),
                                     (HISTORY, history_index, ('society', 'plot', 'charge_type', 'month', 'amount'))):
            if index is not None:
                missing = [h for k, h, _w, _l in (HOUSE_COLS if title == HOUSES else HISTORY_COLS)
                           if k in needed and k not in index]
                if missing:
                    raise UserError(_('Sheet "%(sheet)s" misses the column(s): %(cols)s',
                                      sheet=title, cols=', '.join(missing)))

        lists, phases, sectors = self._lists()
        plan = {
            'errors': [], 'warnings': [], 'stats': collections.Counter(),
            'members_new': collections.OrderedDict(),   # key -> vals
            'members_write': collections.defaultdict(dict),   # member -> vals
            'files_new': collections.OrderedDict(),     # (society id, PLOT) -> vals (+ '_member' key)
            'files_write': collections.defaultdict(dict),     # file -> vals
            'history_new': [], 'history_write': [],
        }
        plan['file_map'] = file_map = {}
        err = lambda sheet, row, msg: plan['errors'].append((sheet, row, msg))
        warn = lambda sheet, row, msg: plan['warnings'].append((sheet, row, msg))

        Member = self.env['res.member'].with_context(active_test=False)
        files = self.env['file'].with_context(active_test=False).search(
            [('society_id', 'in', self.society_ids.ids), ('unit_number', '!=', False)], order='id')
        for f in files:
            key = (f.society_id.id, f.unit_number.strip().upper())
            if key not in file_map or file_map[key].file_status in ('cancel', 'merged_and_cancel'):
                file_map[key] = f
        default_society = self.society_ids if len(self.society_ids) == 1 else self.env['society']

        def society_of(sheet, row, value):
            name = _norm(_text(value))
            if not name:
                if default_society:
                    return default_society
                err(sheet, row, _('Society is empty.'))
                return None
            soc = lists['Society'].get(name)
            if not soc:
                err(sheet, row, _('Society "%s" is not a society of %s.', _text(value), self.company_id.name))
            return soc

        def pick(sheet, row, list_name, value, label, records=None):
            """Record / selection key for a drop-down cell, None when empty or unknown."""
            text = _text(value)
            if not text:
                return None
            if list_name == 'Yes/No':
                if _norm(text) in YES:
                    return True
                if _norm(text) in NO:
                    return False
                warn(sheet, row, _('%(col)s "%(v)s" is not Yes or No: ignored.', col=label, v=text))
                return None
            options = records if records is not None else lists[list_name]
            found = options.get(_norm(text))
            if found is None and list_name in ('Ownership', 'Relation'):
                code = _norm(text).replace(' ', '_')
                found = code if code in options.values() else None
            if found is None:
                warn(sheet, row, _('%(col)s "%(v)s" is not in the list: ignored.', col=label, v=text))
            return found

        # ---------------------------------------------------------- houses
        members_by_ref = {}
        members_by_cnic = collections.defaultdict(list)
        all_members = Member.search([])
        for m in all_members:
            if m.ref:
                members_by_ref[m.ref.strip().upper()] = m
            if m.cnic:
                members_by_cnic[re.sub(r'\D', '', m.cnic)].append(m)
        emails = {(_norm(m.email)): m for m in all_members if m.email}
        new_member_by_ref, new_member_by_cnic = {}, {}
        seen_plots = {}
        plan['house_keys'] = set()

        for row, v in house_rows:
            sheet = HOUSES
            soc = society_of(sheet, row, v.get('society'))
            plot = _text(v.get('plot')).upper()
            if not plot:
                err(sheet, row, _('Plot No is empty.'))
            if not soc or not plot:
                continue
            key = (soc.id, plot)
            if key in seen_plots:
                err(sheet, row, _('Plot %(plot)s is already on row %(row)s.', plot=plot, row=seen_plots[key]))
                continue
            seen_plots[key] = row
            file_rec = file_map.get(key)

            # phase / sector belong to the society of this row
            soc_phases = {_norm(p.name): p for p in phases.sorted('id', reverse=True) if p.society_id == soc}
            phase = pick(sheet, row, 'Phase', v.get('phase'), 'Phase', soc_phases)
            if phase is None and not file_rec and len(soc_phases) == 1:
                phase = next(iter(soc_phases.values()))
            soc_sectors = {_norm(s.name): s for s in sectors.sorted('id', reverse=True) if s.society_id == soc}
            sector = pick(sheet, row, 'Sector', v.get('sector'), 'Sector', soc_sectors)

            # ---- member
            name = _text(v.get('member'))
            ref = _text(v.get('tracking')).upper()
            cnic = _format_cnic(_text(v.get('cnic'))) if _text(v.get('cnic')) else ''
            mvals = {}
            if name:
                mvals['name'] = name
            if cnic:
                mvals['cnic'] = cnic
            mobile_text = _text(v.get('mobile'))
            mobile_digits = re.sub(r'[^1-9]', '', mobile_text)
            if len(mobile_digits) >= 7:   # "0000000000" / "+92" placeholders are skipped
                mobile = _mobile(mobile_text)
                if mobile:
                    mvals['mobile'] = mobile
                else:
                    mvals['_bad_mobile'] = mobile_text
            for col, field in (('phone', 'phone'), ('address', 'street'), ('city', 'city'),
                               ('kin_name', 'kin_name')):
                if _text(v.get(col)):
                    mvals[field] = _text(v.get(col))
            if _text(v.get('kin_cnic')):
                mvals['kin_cnic'] = _format_cnic(_text(v.get('kin_cnic')))
            if _text(v.get('kin_mobile')):
                mvals['kin_mobile'] = _text(v.get('kin_mobile'))
            relation = pick(sheet, row, 'Relation', v.get('kin_relation'), 'Nominee Relation')
            if relation:
                mvals['kin_member_relation'] = relation
            email = _text(v.get('email'))
            if email:
                mvals['email'] = email

            member, member_key = None, None
            if ref and ref in members_by_ref:
                member = members_by_ref[ref]
            elif ref and ref in new_member_by_ref:
                member_key = new_member_by_ref[ref]
            elif not ref and file_rec and file_rec.membership_id:
                member = file_rec.membership_id
            elif cnic:
                digits = re.sub(r'\D', '', cnic)
                member = next((m for m in members_by_cnic.get(digits, []) if _name_key(m.name) == _name_key(name)), None)
                if not member and digits in new_member_by_cnic and _name_key(
                        plan['members_new'][new_member_by_cnic[digits]].get('name')) == _name_key(name):
                    member_key = new_member_by_cnic[digits]
            if member and name and _name_key(member.name) != _name_key(name):
                if file_rec and file_rec.membership_id == member:
                    warn(sheet, row, _('Member name changes from "%(old)s" to "%(new)s".', old=member.name, new=name))
                else:
                    err(sheet, row, _('Tracking ID %(ref)s belongs to "%(old)s", not "%(new)s". '
                                      'Fix the name or empty the Tracking ID for a new member.',
                                      ref=member.ref, old=member.name, new=name))
                    continue
            if 'email' in mvals:
                owner = emails.get(_norm(email))
                if owner and owner != member:
                    warn(sheet, row, _('Email %(e)s is used by member %(m)s: not saved.', e=email, m=owner.name))
                    del mvals['email']
                elif member:
                    emails[_norm(email)] = member
            if '_bad_mobile' in mvals:
                bad = mvals.pop('_bad_mobile')
                if not member or not _same(member.mobile, bad) and not _same(member.phone, bad):
                    warn(sheet, row, _('Mobile "%s" is not like 0300-1234567: saved as Phone.', bad))
                    mvals.setdefault('phone', bad)
            if member:
                changes = {f: val for f, val in mvals.items() if not _same(member[f], val, f)}
                if changes:
                    plan['members_write'][member].update(changes)
            elif member_key:
                plan['members_new'][member_key].update({f: val for f, val in mvals.items()
                                                        if not plan['members_new'][member_key].get(f)})
            else:
                if not name:
                    err(sheet, row, _('Member Name is empty.'))
                    continue
                member_key = 'new-%s' % row
                if ref:
                    mvals['ref'] = ref
                    new_member_by_ref[ref] = member_key
                if cnic:
                    new_member_by_cnic[re.sub(r'\D', '', cnic)] = member_key
                mvals['company_type'] = 'person'
                plan['members_new'][member_key] = mvals

            # ---- file
            fvals = {}
            if phase:
                fvals['phase_id'] = phase.id
            if sector:
                fvals['sector_id'] = sector.id
            for col, field, list_name in (('category', 'category_id', 'Category'),
                                          ('product', 'unit_category_type_id', 'Product'),
                                          ('unit_class', 'unit_class_id', 'Unit Class'),
                                          ('agent', 'maintenance_recovery_agent_id', 'Recovery Agent')):
                rec = pick(sheet, row, list_name, v.get(col), dict((k, h) for k, h, _w, _l in HOUSE_COLS)[col])
                if rec:
                    fvals[field] = rec.id
            ownership = pick(sheet, row, 'Ownership', v.get('ownership'), 'Ownership Type')
            if ownership:
                fvals['file_ownership_type'] = ownership
            for col, field in (('exemption', 'exemption'), ('service_charge', 'service_charge')):
                yes = pick(sheet, row, 'Yes/No', v.get(col), col.replace('_', ' ').title())
                if yes is not None:
                    fvals[field] = yes
            for col, field in (('area', 'standard_area'), ('meter_reading', 'opening_meter_reading')):
                num = _number(v.get(col))
                if num is False:
                    warn(sheet, row, _('%(col)s "%(v)s" is not a number: ignored.', col=col, v=_text(v.get(col))))
                elif num is not None:
                    fvals[field] = num
            booking = _date(v.get('booking_date'))
            if booking is False:
                warn(sheet, row, _('Booking Date "%s" is not a date: ignored.', _text(v.get('booking_date'))))
            elif booking:
                fvals['booking_date'] = booking
            for col, field in (('meter_no', 'meter_no'), ('notes', 'notes')):
                if _text(v.get(col)):
                    fvals[field] = _text(v.get(col))

            if file_rec:
                if _text(v.get('file_no')) and _text(v.get('file_no')).upper() != (file_rec.name or '').upper():
                    warn(sheet, row, _('File No %(x)s differs from Odoo (%(o)s): File No is never changed.',
                                       x=_text(v.get('file_no')), o=file_rec.name))
                if member and member != file_rec.membership_id or member_key:
                    warn(sheet, row, _('Plot %(plot)s changes member from "%(old)s" to "%(new)s".',
                                       plot=plot, old=file_rec.membership_id.name or '-', new=name))
                    fvals['_member'] = member or member_key
                changes = {}
                for f, val in fvals.items():
                    if f == '_member':
                        changes[f] = val
                        continue
                    cur = file_rec[f]
                    cur = cur.id if isinstance(cur, models.BaseModel) else cur
                    if (cur or False) != (val or False):
                        changes[f] = val
                if changes:
                    plan['files_write'][file_rec].update(changes)
                    plan['stats']['houses_updated'] += 1
                else:
                    plan['stats']['houses_unchanged'] += 1
            else:
                fvals.update({'society_id': soc.id, 'unit_number': plot, '_member': member or member_key,
                              'project_type': soc.project_type or 'housing_society', 'type': 'normal'})
                inventory = self.env['plot.inventory'].search(
                    [('society_id', '=', soc.id), ('name', '=ilike', plot), ('file_id', '=', False)], limit=1)
                if inventory:
                    fvals['inventory_id'] = inventory.id
                plan['files_new'][key] = fvals
                plan['stats']['houses_new'] += 1
            plan['house_keys'].add(key)

        plan['stats']['members_new'] = len(plan['members_new'])
        plan['stats']['members_updated'] = len(plan['members_write'])

        # ---------------------------------------------------------- bill history
        file_ids = [f.id for f in file_map.values()]
        existing = collections.defaultdict(list)
        for h in self.env['maintenance.charges.history'].search([('file_id', 'in', file_ids)], order='id'):
            existing[(h.file_id.id, h.charge_type, h.date.replace(day=1))].append(h)
        seen_rows = {}
        for row, v in history_rows:
            sheet = HISTORY
            soc = society_of(sheet, row, v.get('society'))
            plot = _text(v.get('plot')).upper()
            if not plot:
                err(sheet, row, _('Plot No is empty.'))
            if not soc or not plot:
                continue
            key = (soc.id, plot)
            file_rec = file_map.get(key)
            if not file_rec and key not in plan['files_new']:
                err(sheet, row, _('Plot %(plot)s of %(soc)s is not in Odoo and not in sheet Houses.',
                                  plot=plot, soc=soc.name))
                continue
            ctype = CHARGE_TYPES.get(_norm(_text(v.get('charge_type'))))
            month = _month(v.get('month'))
            amount = _number(v.get('amount'))
            paid = _number(v.get('paid'))
            balance = _number(v.get('balance'))
            pay_date = _date(v.get('payment_date'))
            problems = []
            if not ctype:
                problems.append(_('Charge Type must be Utility or Electricity (is "%s").', _text(v.get('charge_type'))))
            if not month:
                problems.append(_('Month "%s" is not a month like Sep-2026.', _text(v.get('month'))))
            if amount is None or amount is False:
                problems.append(_('Bill Amount "%s" is not a number.', _text(v.get('amount'))))
            if paid is False or balance is False:
                problems.append(_('Paid Amount / Balance must be numbers.'))
            if problems:
                for p in problems:
                    err(sheet, row, p)
                continue
            if pay_date is False:
                warn(sheet, row, _('Payment Date "%s" is not a date: ignored.', _text(v.get('payment_date'))))
                pay_date = None
            paid = paid or 0.0
            if balance is None:
                balance = max(amount - paid, 0.0)
            hkey = (key, ctype, month)
            if hkey in seen_rows:
                err(sheet, row, _('%(plot)s %(type)s %(month)s is already on row %(row)s.', plot=plot,
                                  type=CHARGE_LABEL[ctype], month=month.strftime('%b-%Y'), row=seen_rows[hkey]))
                continue
            seen_rows[hkey] = row
            status = 'paid' if balance <= 0.005 and amount > 0 else ('partial' if paid > 0 else 'not_paid')
            vals = {'charge_type': ctype, 'date': month, 'amount': amount, 'amount_paid': paid,
                    'residual': balance, 'payment_status': status, 'payment_date': pay_date or False}
            old = existing.get((file_rec.id, ctype, month), []) if file_rec else []
            if old and any(h.invoice_id for h in old):
                plan['stats']['history_from_bill'] += 1
                continue
            if old:
                h = old[0]
                changes = {f: val for f, val in vals.items()
                           if f in ('amount', 'amount_paid', 'residual') and abs((h[f] or 0.0) - val) > 0.005
                           or f == 'payment_status' and h[f] != val
                           or f == 'payment_date' and (h[f] or False) != val}
                if changes:
                    # payment_date is computed from the amounts: always write it along
                    changes['payment_date'] = vals['payment_date']
                    plan['history_write'].append((h, changes))
                    plan['stats']['history_updated'] += 1
                else:
                    plan['stats']['history_unchanged'] += 1
            else:
                plan['history_new'].append((key, vals))
                plan['stats']['history_new'] += 1
        plan['stats']['house_rows'] = len(house_rows)
        plan['stats']['history_rows'] = len(history_rows)
        return plan

    def _report(self, plan, done=False):
        s = plan['stats']
        head = _('Imported') if done else _('Check result - nothing is saved yet')
        rows = [
            (_('Houses (rows in sheet)'), s['house_rows']),
            (_('New houses (files)'), s['houses_new']),
            (_('Houses updated'), s['houses_updated']),
            (_('Houses unchanged'), s['houses_unchanged']),
            (_('New members'), s['members_new']),
            (_('Members updated'), s['members_updated']),
            (_('Bill rows (rows in sheet)'), s['history_rows']),
            (_('New bill rows'), s['history_new']),
            (_('Bill rows updated'), s['history_updated']),
            (_('Bill rows unchanged'), s['history_unchanged']),
            (_('Rows from real bills (kept as they are)'), s['history_from_bill']),
            (_('Rows with errors (skipped)'), len(plan['errors'])),
            (_('Warnings'), len(plan['warnings'])),
        ]
        html = Markup('<h4 class="o_mei_head">%s</h4><table class="table table-sm table-bordered o_mei_summary" '
                      'style="width:auto">') % head
        for label, value in rows:
            html += Markup('<tr><td>%s</td><td class="text-end" data-count="%s"><b>%s</b></td></tr>') % (
                label, value, value)
        html += Markup('</table>')
        for title, items, cls in ((_('Errors (these rows are skipped)'), plan['errors'], 'text-danger o_mei_errors'),
                                  (_('Warnings'), plan['warnings'], 'text-warning o_mei_warnings')):
            if not items:
                continue
            html += Markup('<h5 class="%s">%s (%s)</h5><table class="table table-sm table-striped %s">'
                           '<tr><th>Sheet</th><th>Row</th><th>Problem</th></tr>') % (cls, title, len(items), cls)
            for sheet, row, msg in items[:500]:
                html += Markup('<tr><td>%s</td><td>%s</td><td>%s</td></tr>') % (sheet, row, msg)
            if len(items) > 500:
                html += Markup('<tr><td colspan="3">%s</td></tr>') % _('... and %s more', len(items) - 500)
            html += Markup('</table>')
        return html

    def action_check(self):
        self.ensure_one()
        plan = self._plan()
        self.write({'state': 'checked', 'result_html': self._report(plan), 'error_count': len(plan['errors'])})
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        plan = self._plan()
        company = self.company_id
        Member = self.env['res.member'].with_company(company)
        File = self.env['file'].with_company(company)

        # members
        new_members = {}
        if plan['members_new']:
            keys = list(plan['members_new'])
            created = Member.create([plan['members_new'][k] for k in keys])
            new_members = dict(zip(keys, created))
        for member, vals in plan['members_write'].items():
            member.with_company(company).write(vals)

        def member_of(token):
            return new_members[token] if isinstance(token, str) else token

        # files
        touched = self.env['file']
        new_files = {}
        if plan['files_new']:
            keys = list(plan['files_new'])
            vals_list = []
            for k in keys:
                vals = dict(plan['files_new'][k])
                member = member_of(vals.pop('_member'))
                vals.update({'membership_id': member.id, 'tracking_id': member.ref})
                vals_list.append(vals)
            created = File.create(vals_list)
            for k, rec, vals in zip(keys, created, vals_list):
                if not vals.get('inventory_id') and rec.unit_number != vals['unit_number']:
                    rec.unit_number = vals['unit_number']
                new_files[k] = rec
            touched |= created
        for file_rec, vals in plan['files_write'].items():
            vals = dict(vals)
            if '_member' in vals:
                member = member_of(vals.pop('_member'))
                vals.update({'membership_id': member.id, 'tracking_id': member.ref})
            file_rec.with_company(company).write(vals)
            touched |= file_rec

        # bill history
        History = self.env['maintenance.charges.history']
        renumber = set()
        new_vals = []
        for key, vals in plan['history_new']:
            file_rec = new_files.get(key) or plan['file_map'][key]
            new_vals.append(dict(vals, file_id=file_rec.id))
            renumber.add((file_rec.id, vals['charge_type']))
            touched |= file_rec
        if new_vals:
            History.create(new_vals)
        for h, vals in plan['history_write']:
            h.write(vals)
            touched |= h.file_id
        for file_id, ctype in renumber:
            rows = History.search([('file_id', '=', file_id), ('charge_type', '=', ctype)], order='date, id')
            for n, h in enumerate(rows, 1):
                if h.installment_number != n:
                    h.installment_number = n

        self.write({'state': 'done', 'result_html': self._report(plan, done=True),
                    'error_count': len(plan['errors']), 'file_ids': [(6, 0, touched.ids)]})
        return self._reopen()

    def action_show_files(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Imported Houses'),
            'res_model': 'file',
            'view_mode': 'list,form',
            'domain': [('id', 'in', self.file_ids.ids)],
            'context': {'active_test': False},
        }

    def action_reset(self):
        self.write({'state': 'draft', 'import_file': False, 'import_name': False, 'result_html': False,
                    'file_ids': [(5,)]})
        return self._reopen()

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            'name': _('Import from Excel'),
        }
