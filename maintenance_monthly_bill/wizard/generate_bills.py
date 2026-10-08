# -*- coding: utf-8 -*-

import base64
import io
import re

from dateutil.relativedelta import relativedelta

from markupsafe import Markup, escape

from odoo import models, fields, api, _
from odoo.exceptions import UserError

DEFAULT_BANK = 'ALLIED BANK LIMITED A/C: 0010046647830014 BRANCH CODE: 0988 (MIDLAND DEVELOPERS PVT LTD)'
DEFAULT_NOTE = ('In case of online payment please use the above bank account and share the '
                'payment screenshot on WhatsApp, or submit the receipt at the society office.')


class MaintenanceBillGenerate(models.TransientModel):
    _name = 'maintenance.bill.generate'
    _description = 'Generate Maintenance Monthly Bills'

    bill_type = fields.Selection([('utility', 'Utility / Maintenance'), ('electricity', 'Electricity')],
                                 required=True, default=lambda self: self.env.context.get('default_bill_type') or 'utility')
    society_id = fields.Many2one('society', required=True, domain="[('is_society', '=', True)]")
    phase_id = fields.Many2one('society', domain="[('is_society', '!=', True), ('society_id', '=', society_id)]")
    unit_numbers = fields.Char(string='House Numbers',
                               help='Plot / house numbers separated by commas, e.g. CV2-R-42, CV2-R-43. '
                                    'Leave empty to bill every file of the society/phase that has a member.')
    file_ids = fields.Many2many('file', string='Only These Houses',
                                domain="[('society_id', '=', society_id), ('phase_id', '=?', phase_id), ('membership_id', '!=', False)]",
                                help='Pick houses of the selected society / phase. Leave empty to bill every house that has a member.')
    bill_month = fields.Date(required=True, default=lambda self: fields.Date.context_today(self).replace(day=1))
    due_date = fields.Date(required=True, default=lambda self: fields.Date.context_today(self).replace(day=1) + relativedelta(days=9))
    journal_id = fields.Many2one('account.journal', required=True, domain="[('type', '=', 'sale')]")

    utility_product_id = fields.Many2one('product.product',
                                         default=lambda self: self._default_product('Maintenance Charges'))
    default_utility_amount = fields.Float(
        string='Utility Charges (no rule)', help='Used only for houses that no Maintenance Charges rule matches.')
    rule_preview = fields.Html(string='Charges per House', compute='_compute_rule_preview', sanitize=False,
                               help='Charge type and amount each house gets from Maintenance Charges (by size).')
    electricity_product_id = fields.Many2one('product.product',
                                             default=lambda self: self._default_product('Electricity'))
    unit_rate = fields.Float(string='Electricity Rate / Unit', digits=(16, 2))
    surcharge_percent = fields.Float(string='Late Surcharge %', default=10.0)
    bank_note = fields.Char(default=DEFAULT_BANK)
    payment_note = fields.Text(default=DEFAULT_NOTE)

    # Electricity readings from Excel (action_download_template / action_import_readings)
    import_file = fields.Binary(string='Readings File (Excel)', attachment=False,
                                help='Excel file made with "Download Template", with the Current Reading filled in.')
    import_filename = fields.Char()
    template_file = fields.Binary(attachment=False)

    def _default_product(self, name):
        if name == 'Maintenance Charges':
            product = self.env['maintenance.charges']._get_maintenance_charges_product_id()
            if product:
                return product
        return self.env['product.product'].search([('name', '=', name), ('type', '=', 'service')], limit=1)

    @api.onchange('society_id')
    def _onchange_society(self):
        """Default to the sales journal of the company named like the society (CV-I -> Canal Valley-I)."""
        if self.society_id:
            company = self.env['res.company'].search([('name', '=', self.society_id.name)], limit=1) \
                or self.society_id.company_id
            self.journal_id = self.env['account.journal'].search(
                [('type', '=', 'sale'), ('company_id', '=', company.id)], limit=1)
            self.phase_id = False
            self.file_ids = False
            if not self.journal_id:
                return {'warning': {
                    'title': _('Company not selected'),
                    'message': _('No sales journal found for %s. Tick that company in the company '
                                 'switcher (top right), then choose the society again.') % company.name,
                }}

    @api.onchange('phase_id')
    def _onchange_phase(self):
        # keep only houses that belong to the newly chosen phase
        if self.phase_id:
            self.file_ids = self.file_ids.filtered(lambda f: f.phase_id == self.phase_id)

    def _utility_rules(self):
        """Maintenance Charges rule lines of the society / phase in force during the billed month."""
        month = self.bill_month.replace(day=1)
        month_end = month + relativedelta(months=1, days=-1)
        domain = [
            ('maintenance_charges_id.society_id', '=', self.society_id.id),
            ('maintenance_charges_id.date_from', '<=', month_end),
            ('maintenance_charges_id.date_to', '>=', month),
            ('maintenance_charges_type_id', '!=', False),
        ]
        if self.phase_id:
            domain.append(('maintenance_charges_id.phase_id', '=', self.phase_id.id))
        return self.env['maintenance.charges.line'].search(domain)

    @staticmethod
    def _house_marla(file_rec):
        """Size of the house in marla: Area (Marla) of its size, or the number in the size name
        ("10 Marla") when that area is not filled in."""
        size = file_rec.unit_category_type_id
        if size.area_marla:
            return size.area_marla
        found = re.search(r'(\d+(?:\.\d+)?)\s*marla', size.name or '', re.IGNORECASE)
        return float(found.group(1)) if found else 0.0

    def _utility_rule_for(self, file_rec, rules=None):
        """(charge type, charge type line) of the house, from Maintenance Charges: same phase,
        sector (when the rule lists sectors), category, type (Plot / House / Shop) and the house
        size in marla within From / To (3.5 marla counts as 3). When ranges overlap (1-3, 1-4,
        1-5) the narrowest one wins, so a 3 marla house gets the 3 marla charge."""
        rules = self._utility_rules() if rules is None else rules
        area = self._house_marla(file_rec)
        matching = rules.filtered(lambda l: (
            l.maintenance_charges_id.phase_id == file_rec.phase_id
            and (not l.maintenance_charges_id.sector_ids or file_rec.sector_id in l.maintenance_charges_id.sector_ids)
            and l.category_id == file_rec.category_id
            and (not l.unit_class_id or l.unit_class_id == file_rec.unit_class_id)
            and l.from_no <= area < l.to_no + 1))
        for line in matching.sorted(lambda l: (l.to_no - l.from_no, -l.maintenance_charges_id.date_from.toordinal(), -l.id)):
            charges = line.maintenance_charges_type_id.maintenance_charges_type_line_ids
            charge = charges.filtered(lambda c: c.product_id == self.utility_product_id)[:1] or charges[:1]
            if charge:
                return line.maintenance_charges_type_id, charge
        return self.env['maintenance.charges.type'], self.env['maintenance.charges.type.lines']

    def _utility_values_for(self, file_rec, rules=None):
        """Charge type, product and amount (after exemption) of the house's utility bill."""
        charge_type, charge = self._utility_rule_for(file_rec, rules)
        product = charge.product_id or self.utility_product_id
        amount = charge.amount if charge else self.default_utility_amount
        return charge_type, product, self._apply_exemption(file_rec, amount, product)

    @api.depends('bill_type', 'society_id', 'phase_id', 'file_ids', 'bill_month', 'utility_product_id',
                 'default_utility_amount')
    def _compute_rule_preview(self):
        for wizard in self:
            wizard.rule_preview = False
            if wizard.bill_type != 'utility' or not wizard.society_id or not wizard.bill_month:
                continue
            if wizard.file_ids:
                files = wizard.file_ids
            else:
                domain = [('society_id', '=', wizard.society_id.id), ('membership_id', '!=', False)]
                if wizard.phase_id:
                    domain.append(('phase_id', '=', wizard.phase_id.id))
                files = self.env['file'].search(domain)
            rules = wizard._utility_rules()
            groups = {}
            for file_rec in files:
                charge_type, charge = wizard._utility_rule_for(file_rec, rules)
                key = (charge_type.name or _('No rule found: Utility Charges below'),
                       charge.amount if charge else wizard.default_utility_amount)
                groups.setdefault(key, []).append(file_rec.unit_number or file_rec.name or '')
            rows = ''.join(
                '<tr><td>%s</td><td class="text-end">%s</td><td class="text-end">%s</td><td class="text-muted">%s</td></tr>' % (
                    escape(name), len(units), escape('{:,.0f}'.format(amount)),
                    escape(', '.join(sorted(units)[:6]) + (' ...' if len(units) > 6 else '')))
                for (name, amount), units in sorted(groups.items()))
            wizard.rule_preview = Markup(
                '<table class="table table-sm mb-0"><thead><tr><th>Charge Type</th><th class="text-end">Houses</th>'
                '<th class="text-end">Amount</th><th>e.g.</th></tr></thead><tbody>%s</tbody></table>' % rows
            ) if rows else False

    def _apply_exemption(self, file_rec, amount, product=None):
        """Approved Maintenance Exemption of the file that covers the billed month."""
        if 'maintenance.exemption.history' not in self.env:
            return amount
        month_end = self.bill_month.replace(day=1) + relativedelta(months=1, days=-1)
        exemption = self.env['maintenance.exemption.history'].search([
            ('file_id', '=', file_rec.id), ('exemption_state', '=', 'active'),
            ('product_id', '=', (product or self.utility_product_id).id),
            ('from_date', '<=', month_end), ('to_date', '>=', self.bill_month.replace(day=1)),
        ], limit=1)
        if not exemption:
            return amount
        if exemption.exemption_nature == 'full':
            return 0.0
        if exemption.exemption_type == 'percentage' and exemption.exemption_percent:
            return round(amount * (100 - exemption.exemption_percent) / 100.0, 2)
        if exemption.exemption_type == 'fixed_amount' and exemption.exemption_amount:
            return max(amount - exemption.exemption_amount, 0.0)
        return amount

    def _bill_vals(self, file_rec, month, draft=None):
        """Values of a new bill of file_rec, or the values to refresh an existing draft with."""
        Bill = self.env['maintenance.bill']
        vals = {
            'journal_id': self.journal_id.id,
            'due_date': self.due_date,
            'surcharge_percent': self.surcharge_percent,
            'bank_note': self.bank_note,
            'payment_note': self.payment_note,
        }
        if self.bill_type == 'utility':
            charge_type, product, amount = self._utility_values_for(file_rec)
            vals['maintenance_charges_type_id'] = charge_type.id
            vals['utility_product_id'] = product.id
            # an amount typed on the draft is kept when nothing better is known
            if not draft or amount:
                vals['utility_amount'] = amount
        else:
            vals.update({'electricity_product_id': self.electricity_product_id.id, 'unit_rate': self.unit_rate})
            # the meter continues from the last electricity reading of this house; a draft that
            # already has a current reading keeps its readings
            if not draft or not draft.current_reading:
                last = Bill.search([('file_id', '=', file_rec.id), ('bill_month', '<', month), ('state', '!=', 'cancel'),
                                    ('bill_type', 'in', ('electricity', 'combined'))], order='bill_month desc', limit=1)
                vals.update({
                    'meter_no': last.meter_no or file_rec.meter_no,
                    'previous_reading': last.current_reading if last else file_rec.opening_meter_reading,
                })
        return vals

    def _selected_files(self):
        """Houses chosen in the wizard: picked houses, typed house numbers, or the whole society / phase."""
        if self.file_ids:
            files = self.file_ids
        elif self.unit_numbers:
            wanted = [u.strip().upper() for u in self.unit_numbers.replace('\n', ',').split(',') if u.strip()]
            files = self.env['file'].search([('society_id', '=', self.society_id.id),
                                             ('unit_number', 'in', wanted), ('membership_id', '!=', False)])
            missing = set(wanted) - {(u or '').upper() for u in files.mapped('unit_number')}
            if missing:
                raise UserError(_('No file with a member found in %(society)s for: %(units)s',
                                  society=self.society_id.name, units=', '.join(sorted(missing))))
        else:
            domain = [('society_id', '=', self.society_id.id), ('membership_id', '!=', False)]
            if self.phase_id:
                domain.append(('phase_id', '=', self.phase_id.id))
            files = self.env['file'].search(domain)
        if not files:
            raise UserError(_('No files with a member found for this selection.'))
        return files

    def _check_settings(self):
        if self.bill_type == 'utility' and not self.utility_product_id:
            raise UserError(_('Set the Utility product first.'))
        if self.bill_type == 'electricity' and not self.electricity_product_id:
            raise UserError(_('Set the Electricity product first.'))
        # 2026-10-06: drafts were generated with "Maintenance Charges" as the electricity product
        maintenance_product = self._default_product('Maintenance Charges')
        if self.bill_type == 'electricity' and self.electricity_product_id == maintenance_product:
            raise UserError(_('Electricity Product is "%s", the utility product. Pick the Electricity product.',
                              self.electricity_product_id.display_name))
        if self.bill_type == 'utility' and self.utility_product_id == self._default_product('Electricity'):
            raise UserError(_('Utility Product is "%s", the electricity product. Pick the Maintenance Charges product.',
                              self.utility_product_id.display_name))
        if self.journal_id.type != 'sale':
            raise UserError(_('Pick a Sales journal.'))

    def _generate_bills(self, files, overrides=None):
        """Create this month's bills of this type for files, or refresh their drafts.
        overrides = {file id: values} written over the generated values (readings from Excel).
        Returns (new bills, refreshed drafts, existing bills of the month)."""
        self._check_settings()
        overrides = overrides or {}
        month = self.bill_month.replace(day=1)
        Bill = self.env['maintenance.bill']
        # a house gets one bill of this type per month (an old combined bill counts as both).
        # Posted bills are kept; draft bills of this type are refreshed with the values entered
        # here (rate, charges, dates), so generating again after a mistake fixes the drafts.
        existing = Bill.search([('file_id', 'in', files.ids), ('bill_month', '=', month), ('state', '!=', 'cancel'),
                                ('bill_type', 'in', (self.bill_type, 'combined'))])
        drafts = existing.filtered(lambda b: b.state == 'draft' and b.bill_type == self.bill_type)
        vals_list = []
        for file_rec in files - existing.mapped('file_id'):
            vals = dict(self._bill_vals(file_rec, month), bill_type=self.bill_type, file_id=file_rec.id, bill_month=month)
            vals.update(overrides.get(file_rec.id, {}))
            vals_list.append(vals)
        bills = Bill.create(vals_list)
        for bill in drafts:
            bill.write(dict(self._bill_vals(bill.file_id, month, bill), **overrides.get(bill.file_id.id, {})))
        (bills | drafts).action_recompute_arrears()
        return bills, drafts, existing

    def _bills_action(self, bills, drafts, existing, note=''):
        month = self.bill_month.replace(day=1)
        kept = existing - drafts
        kind = dict(self._fields['bill_type']._description_selection(self.env))[self.bill_type]
        return {
            'type': 'ir.actions.act_window',
            'name': _('%(kind)s Bills %(month)s (%(new)s new, %(updated)s drafts updated, %(skip)s already posted)%(note)s',
                      kind=kind, month=month.strftime('%b %Y'), new=len(bills), updated=len(drafts), skip=len(kept),
                      note=note),
            'res_model': 'maintenance.bill',
            'view_mode': 'list,form',
            'domain': [('id', 'in', (bills | existing).ids)],
            'views': [(self.env.ref('maintenance_monthly_bill.maintenance_bill_list_%s' % self.bill_type).id, 'list'),
                      (False, 'form')],
            'context': {'search_default_group_state': 0, 'default_bill_type': self.bill_type},
        }

    def action_generate(self):
        self.ensure_one()
        bills, drafts, existing = self._generate_bills(self._selected_files())
        return self._bills_action(bills, drafts, existing)

    # ------------------------------------------------------------------ electricity readings from Excel
    # (column header, key). House No and Current Reading are required; an empty Meter No, Previous
    # Reading or Rate / Unit keeps the value the wizard would use (last bill / wizard rate).
    READING_COLUMNS = [
        ('House No', 'house'),
        ('Member Name', 'member'),
        ('Meter No', 'meter_no'),
        ('Previous Reading', 'previous_reading'),
        ('Current Reading', 'current_reading'),
        ('Rate / Unit', 'unit_rate'),
    ]

    def _previous_reading(self, file_rec, month):
        last = self.env['maintenance.bill'].search([
            ('file_id', '=', file_rec.id), ('bill_month', '<', month), ('state', '!=', 'cancel'),
            ('bill_type', 'in', ('electricity', 'combined'))], order='bill_month desc', limit=1)
        return (last.meter_no or file_rec.meter_no or '',
                last.current_reading if last else (file_rec.opening_meter_reading or 0.0))

    def action_download_template(self):
        """Excel sheet of the selected houses with their meter and previous reading filled in:
        only the Current Reading column is left to fill."""
        self.ensure_one()
        import xlsxwriter
        month = self.bill_month.replace(day=1)
        files = self._selected_files().sorted(lambda f: f.unit_number or '')
        buffer = io.BytesIO()
        book = xlsxwriter.Workbook(buffer, {'in_memory': True})
        sheet = book.add_worksheet('Electricity Readings')
        head = book.add_format({'bold': True, 'bg_color': '#E8EEF8', 'border': 1, 'border_color': '#C5CEE0'})
        need = book.add_format({'bold': True, 'bg_color': '#FFF2CC', 'border': 1, 'border_color': '#C5CEE0'})
        locked = book.add_format({'bg_color': '#F6F8FB', 'border': 1, 'border_color': '#D9DEE8'})
        cell = book.add_format({'border': 1, 'border_color': '#D9DEE8'})
        for col, (label, key) in enumerate(self.READING_COLUMNS):
            sheet.write(0, col, label + (' *' if key in ('house', 'current_reading') else ''),
                        need if key == 'current_reading' else head)
        for row, file_rec in enumerate(files, start=1):
            meter_no, previous = self._previous_reading(file_rec, month)
            values = [file_rec.unit_number or '', file_rec.membership_id.name or '', meter_no, previous, None,
                      self.unit_rate or None]
            for col, value in enumerate(values):
                fmt = cell if col in (2, 3, 4, 5) else locked
                if value is None:
                    sheet.write_blank(row, col, None, fmt)
                else:
                    sheet.write(row, col, value, fmt)
        for col, width in enumerate((16, 30, 14, 17, 17, 12)):
            sheet.set_column(col, col, width)
        sheet.freeze_panes(1, 0)
        notes = book.add_worksheet('How to fill')
        for r, line in enumerate([
            'Electricity readings for %s, %s' % (self.society_id.name, month.strftime('%B %Y')),
            '',
            'House No *        : plot / house number exactly as in Odoo (e.g. CV-R-126). Do not change it.',
            'Member Name       : for reference only, not imported.',
            'Meter No          : leave as is; change it only if the meter was replaced.',
            "Previous Reading  : last month's reading; leave empty to take it from the last bill.",
            "Current Reading * : this month's meter reading (must not be lower than the previous reading).",
            'Rate / Unit       : leave empty to use the rate entered in Generate Electricity Bills.',
            '',
            'Rows without a Current Reading are skipped. Then: Generate Electricity Bills > Import & Generate.',
        ]):
            notes.write(r, 0, line)
        notes.set_column(0, 0, 110)
        book.close()
        self.template_file = base64.b64encode(buffer.getvalue())
        name = 'Electricity Readings %s %s.xlsx' % (self.society_id.name, month.strftime('%b-%Y'))
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=%s&id=%s&field=template_file&download=true&filename=%s' % (self._name, self.id, name),
            'target': 'self',
        }

    @staticmethod
    def _number(value):
        if value in (None, ''):
            return None
        if isinstance(value, (int, float)):
            return float(value)
        return float(str(value).replace(',', '').strip())

    def _read_readings_file(self):
        """Rows of the uploaded sheet: [(excel row number, {key: value})]."""
        import openpyxl
        if not self.import_file:
            raise UserError(_('Choose the readings Excel file first (make it with "Download Template").'))
        if self.import_filename and not self.import_filename.lower().endswith(('.xlsx', '.xlsm')):
            raise UserError(_('The readings file must be an Excel .xlsx file (use "Download Template").'))
        try:
            book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.import_file)), data_only=True, read_only=True)
        except Exception:
            raise UserError(_('This file could not be read as an Excel .xlsx file. Use "Download Template".'))
        sheet = book.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise UserError(_('The readings file is empty.'))
        wanted = {label.lower(): key for label, key in self.READING_COLUMNS}
        header = [str(h or '').replace('*', '').strip().lower() for h in rows[0]]
        index = {wanted[h]: i for i, h in enumerate(header) if h in wanted}
        missing = [label for label, key in self.READING_COLUMNS if key in ('house', 'current_reading') and key not in index]
        if missing:
            raise UserError(_('Column(s) %s not found in the first row. Use "Download Template".') % ', '.join(missing))
        result = []
        for number, row in enumerate(rows[1:], start=2):
            if not any(v not in (None, '') for v in row):
                continue
            result.append((number, {key: (row[i] if i < len(row) else None) for key, i in index.items()}))
        return result

    def action_import_readings(self):
        """Create / refresh the electricity draft bills from the readings sheet. Every row is
        checked first: if any row is wrong nothing is imported and all errors are listed."""
        self.ensure_one()
        if self.bill_type != 'electricity':
            raise UserError(_('Readings can only be imported for electricity bills.'))
        month = self.bill_month.replace(day=1)
        rows = self._read_readings_file()
        domain = [('society_id', '=', self.society_id.id)]
        if self.phase_id:
            domain.append(('phase_id', '=', self.phase_id.id))
        houses = {}
        for file_rec in self.env['file'].search(domain):
            key = (file_rec.unit_number or '').strip().upper()
            # a house can have old files without a member: prefer the one with a member
            if key and (key not in houses or (file_rec.membership_id and not houses[key].membership_id)):
                houses[key] = file_rec

        errors, overrides, seen, skipped = [], {}, {}, 0
        for number, row in rows:
            house = str(row.get('house') or '').strip()
            if not house:
                errors.append(_('Row %s: House No is empty.', number)); continue
            file_rec = houses.get(house.upper())
            if not file_rec:
                errors.append(_('Row %(row)s: house %(house)s not found in %(society)s.',
                                row=number, house=house, society=self.phase_id.name or self.society_id.name)); continue
            if not file_rec.membership_id:
                errors.append(_('Row %(row)s: house %(house)s has no member.', row=number, house=house)); continue
            if house.upper() in seen:
                errors.append(_('Row %(row)s: house %(house)s is also on row %(first)s.',
                                row=number, house=house, first=seen[house.upper()])); continue
            seen[house.upper()] = number
            try:
                current = self._number(row.get('current_reading'))
                previous = self._number(row.get('previous_reading'))
                rate = self._number(row.get('unit_rate'))
            except ValueError:
                errors.append(_('Row %(row)s (%(house)s): readings and rate must be numbers.', row=number, house=house)); continue
            if current is None:
                skipped += 1      # no reading yet for this house
                continue
            if previous is None:
                previous = self._previous_reading(file_rec, month)[1]
            if current < previous:
                errors.append(_('Row %(row)s (%(house)s): current reading %(cur)s is lower than previous reading %(prev)s.',
                                row=number, house=house, cur=int(current), prev=int(previous))); continue
            vals = {'previous_reading': previous, 'current_reading': current}
            meter = row.get('meter_no')
            if meter not in (None, ''):
                vals['meter_no'] = str(int(meter)) if isinstance(meter, float) and meter.is_integer() else str(meter).strip()
            if rate is not None:
                vals['unit_rate'] = rate
            overrides[file_rec.id] = vals
        if errors:
            shown = errors[:40]
            more = len(errors) - len(shown)
            raise UserError(_('Nothing was imported. Fix these rows and import again:\n\n%s', '\n'.join(shown))
                            + (_('\n... and %s more.', more) if more else ''))
        if not overrides:
            raise UserError(_('No row has a Current Reading: nothing to import.'))
        files = self.env['file'].browse(list(overrides))
        bills, drafts, existing = self._generate_bills(files, overrides)
        note = _(', %(rows)s rows imported', rows=len(overrides))
        if skipped:
            note += _(', %(skip)s rows without reading skipped', skip=skipped)
        return self._bills_action(bills, drafts, existing, note=note)
